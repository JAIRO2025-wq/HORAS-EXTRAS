"""
Genera el informe de OCTUBRE 2026 (horas extra) a partir de la lectura ARREGLADA
de data/records.  Uso:

    <python> build-informe-octubre.py <datos.json> <salida.xlsx>

El JSON de entrada lo produce la lectura corregida de data/records (misma lógica
que /api/admin/all-records ya sin el bug de deduplicación entre empleados).
"""
import json
import sys
from datetime import datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

JSON_PATH = sys.argv[1] if len(sys.argv) > 1 else "datos-informe.json"
OUT_PATH = sys.argv[2] if len(sys.argv) > 2 else "Informe_Octubre_2026.xlsx"
MES = "Octubre"
LOCAL = timezone(timedelta(hours=-6))  # Honduras

AZUL = PatternFill("solid", fgColor="FF0070C0")
VERDE = PatternFill("solid", fgColor="FFC6EFCE")
AMBAR = PatternFill("solid", fgColor="FFFFC000")
GRIS = PatternFill("solid", fgColor="FFF2F2F2")
ROJO = PatternFill("solid", fgColor="FFFFC7CE")
BLANCO_BOLD = Font(color="FFFFFFFF", bold=True)
BOLD = Font(bold=True)
THIN = Side(style="thin", color="FFD9D9D9")
BORDE = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

ESTADO = {"approved": "Aprobado", "pending": "Pendiente", "rejected": "Rechazado"}
TIPO = {"overtime": "Horas Extra", "additional_day": "Día Adicional"}
DIAS = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]

data = json.load(open(JSON_PATH, encoding="utf-8"))
octubre = data["months"][MES]
records = octubre["finales"]
ocultos = octubre["ocultos"]


def fecha(iso):
    return str(iso)[:10]


def dia_semana(iso):
    y, m, d = (int(x) for x in fecha(iso).split("-"))
    return DIAS[datetime(y, m, d).weekday()]


def local_dt(iso):
    if not iso:
        return ""
    try:
        return datetime.fromisoformat(str(iso).replace("Z", "+00:00")).astimezone(LOCAL).strftime("%d/%m/%y %H:%M")
    except Exception:
        return ""


def nombre(r):
    return r.get("officialName") or r.get("employeeName") or "SIN NOMBRE"


def horas(r):
    return float(r.get("dayHours") or 0), float(r.get("nightHours") or 0)


def tarifas(r):
    salario = float(r.get("salary") or 0)
    base = salario / 2 / 15 / 8 if salario > 0 else 0.0
    return salario, base, base * 2, base * 2 * 1.25


# ---------------------------------------------------------------- nómina
aprobados = [r for r in records if r.get("status") == "approved"]
por_empleado = {}
for r in aprobados:
    key = nombre(r)
    e = por_empleado.setdefault(key, {
        "sucursal": r.get("branch") or "—", "salario": 0.0,
        "dh": 0.0, "nh": 0.0, "adic": 0, "acts": [],
    })
    d, n = horas(r)
    e["dh"] += d
    e["nh"] += n
    if r.get("type") == "additional_day":
        e["adic"] += 1
    e["salario"] = float(r.get("salary") or 0)
    if r.get("activity"):
        e["acts"].append(str(r["activity"]).strip())

nomina = []
for emp, e in por_empleado.items():
    base = e["salario"] / 2 / 15 / 8 if e["salario"] > 0 else 0.0
    pago_d = e["dh"] * base * 2
    pago_n = e["nh"] * base * 2 * 1.25
    pago_a = e["adic"] * base * 8
    acts = list(dict.fromkeys(e["acts"]))
    nomina.append([emp, e["sucursal"], e["salario"], base, e["dh"], pago_d, e["nh"], pago_n,
                   e["adic"], pago_a, pago_d + pago_n + pago_a,
                   f"{e['dh'] + e['nh']:.2f} HE. ACTIVIDADES: " + ", ".join(acts)])
nomina.sort(key=lambda x: (x[1], x[0]))

wb = Workbook()

# ---- Hoja 1: Nómina octubre
ws = wb.active
ws.title = "Nómina Octubre"
ws["A1"] = f"RESUMEN DE NÓMINA - OCTUBRE 2026 (solo registros APROBADOS)"
ws["A1"].font = Font(size=14, bold=True)
ws["A2"] = f"Generado el {datetime.now(LOCAL).strftime('%d/%m/%Y %H:%M')} | Fuente: data/records/2026/OCTUBRE (lectura corregida)"
ws["A2"].font = Font(size=9, italic=True, color="FF666666")

cols = ["Empleado", "Sucursal", "Salario Base", "Valor Hora", "H. Diurnas", "Pago Diurno",
        "H. Nocturnas", "Pago Nocturno", "Días Adic.", "Pago Adic.", "TOTAL A PAGAR", "Detalle de Actividades"]
ws.append([])
ws.append(cols)
fila_hdr = ws.max_row
for c in range(1, len(cols) + 1):
    cell = ws.cell(row=fila_hdr, column=c)
    cell.fill = AZUL
    cell.font = BLANCO_BOLD
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

for row in nomina:
    ws.append(row)
    r = ws.max_row
    ws.cell(row=r, column=6).fill = VERDE
    ws.cell(row=r, column=8).fill = VERDE
    ws.cell(row=r, column=12).alignment = Alignment(wrap_text=True, vertical="top")
    for c in (3, 4, 6, 8, 10, 11):
        ws.cell(row=r, column=c).number_format = '"$"#,##0.00'
    for c in (5, 7):
        ws.cell(row=r, column=c).number_format = "0.00"
    ws.cell(row=r, column=11).font = BOLD

if nomina:
    tot = ["TOTALES (%d personas)" % len(nomina), "", "", "", sum(x[4] for x in nomina),
           sum(x[5] for x in nomina), sum(x[6] for x in nomina), sum(x[7] for x in nomina),
           sum(x[8] for x in nomina), sum(x[9] for x in nomina), sum(x[10] for x in nomina), ""]
    ws.append(tot)
    r = ws.max_row
    for c in range(1, 13):
        ws.cell(row=r, column=c).fill = GRIS
        ws.cell(row=r, column=c).font = BOLD
    for c in (3, 4, 6, 8, 10, 11):
        ws.cell(row=r, column=c).number_format = '"$"#,##0.00'
    for c in (5, 7):
        ws.cell(row=r, column=c).number_format = "0.00"

for i, w in enumerate([30, 14, 12, 11, 10, 12, 11, 12, 10, 12, 15, 60], start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = f"A{fila_hdr + 1}"

# ---- Hoja 2: Detalle de octubre
ws2 = wb.create_sheet("Detalle Octubre")
ws2["A1"] = "DETALLE DE REGISTROS - OCTUBRE 2026 (todos los estados, quincenas 1 y 2)"
ws2["A1"].font = Font(size=14, bold=True)
ws2["A2"] = f"{len(records)} registros leídos de data/records/2026/OCTUBRE (antes la app mostraba solo {octubre['mostradosAntes']})"
ws2["A2"].font = Font(size=9, italic=True, color="FF666666")
cols2 = ["Empleado", "Sucursal", "Fecha laboral", "Día", "Horario", "Actividad", "Compañeros",
         "H. Diurnas", "H. Nocturnas", "H. Totales", "Q", "Tipo", "Estado", "Registrado",
         "¿La app lo ocultaba?"]
ws2.append([])
ws2.append(cols2)
hdr2 = ws2.max_row
for c in range(1, len(cols2) + 1):
    cell = ws2.cell(row=hdr2, column=c)
    cell.fill = AZUL
    cell.font = BLANCO_BOLD
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

ocultos_ids = {h.get("id") for h in ocultos}
for r in sorted(records, key=lambda x: (nombre(x), fecha(x["date"]))):
    d, n = horas(r)
    es_oculto = r.get("id") in ocultos_ids
    ws2.append([nombre(r), r.get("branch") or "—", fecha(r["date"]), dia_semana(r["date"]),
                f"{r.get('startTime', '')} - {r.get('endTime', '')}", r.get("activity", ""),
                r.get("coworkers", ""), d, n, float(r.get("totalHours") or 0), r.get("quincena"),
                TIPO.get(r.get("type"), r.get("type")), ESTADO.get(r.get("status"), r.get("status")),
                local_dt(r.get("createdAt")), "SÍ - la ocultaba el bug" if es_oculto else ""])
    fila = ws2.max_row
    for c in (8, 9, 10):
        ws2.cell(row=fila, column=c).number_format = "0.00"
    ws2.cell(row=fila, column=6).alignment = Alignment(wrap_text=True, vertical="top")
    if es_oculto:
        for c in range(1, len(cols2) + 1):
            ws2.cell(row=fila, column=c).fill = ROJO

for i, w in enumerate([30, 14, 13, 11, 20, 52, 28, 10, 11, 10, 5, 13, 11, 15, 22], start=1):
    ws2.column_dimensions[get_column_letter(i)].width = w
ws2.freeze_panes = f"A{hdr2 + 1}"

# ---- Hoja 3: quiénes NO aparecían
ws3 = wb.create_sheet("Ocultos por el bug")
ws3["A1"] = "EMPLEADOS QUE LA APP NO MOSTRABA (OCTUBRE 2026)"
ws3["A1"].font = Font(size=14, bold=True)
ws3["A2"] = ("El anti-duplicados comparaba fecha+horario+actividad sin mirar al empleado: cuando dos personas "
             "que trabajaron juntas registraban lo mismo, el registro de una de ellas se descartaba al leer. "
             "Nadie cobró de más: el registro que sobrevivía era el propio del otro empleado, así que estas horas "
             "simplemente no aparecían en ninguna parte de la planilla.")
ws3["A2"].font = Font(size=9, italic=True, color="FF666666")

resumen = {}
for h in ocultos:
    e = resumen.setdefault(nombre(h), {"sucursal": h.get("branch") or "—", "n": 0, "h": 0.0, "fechas": [], "abs": set()})
    e["n"] += 1
    e["h"] += float(h.get("totalHours") or 0)
    e["fechas"].append(fecha(h["date"]))
    e["abs"].add(h.get("absorbidoPor") or "?")

ws3.append([])
ws3.append(["Empleado", "Sucursal", "Registros ocultos", "Horas ocultas", "Fechas", "Registro que sí aparecía (de otro empleado)"])
hdr3 = ws3.max_row
for c in range(1, 7):
    cell = ws3.cell(row=hdr3, column=c)
    cell.fill = AMBAR
    cell.font = BOLD
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for emp, e in sorted(resumen.items(), key=lambda kv: -kv[1]["h"]):
    ws3.append([emp, e["sucursal"], e["n"], e["h"], ", ".join(sorted(e["fechas"])), ", ".join(sorted(e["abs"]))])
    r = ws3.max_row
    ws3.cell(row=r, column=4).number_format = "0.00"
    ws3.cell(row=r, column=4).font = BOLD
    ws3.cell(row=r, column=6).alignment = Alignment(wrap_text=True, vertical="top")
ws3.append(["TOTAL", "", sum(e["n"] for e in resumen.values()), sum(e["h"] for e in resumen.values()), "", ""])
r = ws3.max_row
for c in range(1, 7):
    ws3.cell(row=r, column=c).fill = GRIS
    ws3.cell(row=r, column=c).font = BOLD
ws3.cell(row=r, column=4).number_format = "0.00"

ws3.append([])
ws3.append(["DETALLE DE LOS REGISTROS OCULTOS"])
ws3.cell(row=ws3.max_row, column=1).font = Font(bold=True)
ws3.append(["Empleado", "Fecha laboral", "Horario", "Actividad", "H. Totales", "Tipo", "Estado", "El registro que aparecía era de"])
hdr3b = ws3.max_row
for c in range(1, 9):
    cell = ws3.cell(row=hdr3b, column=c)
    cell.fill = AZUL
    cell.font = BLANCO_BOLD
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
for h in sorted(ocultos, key=lambda x: (nombre(x), fecha(x["date"]))):
    ws3.append([nombre(h), fecha(h["date"]), f"{h.get('startTime', '')} - {h.get('endTime', '')}",
                h.get("activity", ""), float(h.get("totalHours") or 0), TIPO.get(h.get("type"), ""),
                ESTADO.get(h.get("status"), ""), h.get("absorbidoPor") or "?"])
    ws3.cell(row=ws3.max_row, column=5).number_format = "0.00"
    ws3.cell(row=ws3.max_row, column=4).alignment = Alignment(wrap_text=True, vertical="top")

for i, w in enumerate([30, 14, 16, 13, 34, 30], start=1):
    ws3.column_dimensions[get_column_letter(i)].width = w
for i, w in enumerate([30, 13, 20, 50, 10, 13, 11, 34], start=1):
    letra = get_column_letter(i)
    if ws3.column_dimensions[letra].width < w:
        ws3.column_dimensions[letra].width = w

# ---- Hoja 4: por sucursal
ws4 = wb.create_sheet("Por Sucursal")
ws4["A1"] = "RESUMEN POR SUCURSAL - OCTUBRE 2026 (registros aprobados)"
ws4["A1"].font = Font(size=14, bold=True)
ws4.append([])
ws4.append(["Sucursal", "Empleados", "H. Diurnas", "H. Nocturnas", "Total HE", "Monto a pagar"])
hdr4 = ws4.max_row
for c in range(1, 7):
    cell = ws4.cell(row=hdr4, column=c)
    cell.fill = AMBAR
    cell.font = BOLD
    cell.alignment = Alignment(horizontal="center", wrap_text=True)
suc = {}
for row in nomina:
    s = suc.setdefault(row[1], {"emp": 0, "dh": 0.0, "nh": 0.0, "monto": 0.0})
    s["emp"] += 1
    s["dh"] += row[4]
    s["nh"] += row[6]
    s["monto"] += row[10]
for s, v in sorted(suc.items()):
    ws4.append([s, v["emp"], v["dh"], v["nh"], v["dh"] + v["nh"], v["monto"]])
    r = ws4.max_row
    for c in (3, 4, 5):
        ws4.cell(row=r, column=c).number_format = "0.00"
    ws4.cell(row=r, column=6).number_format = '"$"#,##0.00'
ws4.append(["TOTAL", sum(v["emp"] for v in suc.values()), sum(v["dh"] for v in suc.values()),
            sum(v["nh"] for v in suc.values()), sum(v["dh"] + v["nh"] for v in suc.values()),
            sum(v["monto"] for v in suc.values())])
r = ws4.max_row
for c in range(1, 7):
    ws4.cell(row=r, column=c).fill = GRIS
    ws4.cell(row=r, column=c).font = BOLD
for c in (3, 4, 5):
    ws4.cell(row=r, column=c).number_format = "0.00"
ws4.cell(row=r, column=6).number_format = '"$"#,##0.00'
for i, w in enumerate([18, 12, 12, 13, 11, 15], start=1):
    ws4.column_dimensions[get_column_letter(i)].width = w

# ---- Hoja 5: alcance del bug por mes
ws5 = wb.create_sheet("Bug por mes")
ws5["A1"] = "ALCANCE DEL BUG DEL ANTI-DUPLICADOS (año 2026)"
ws5["A1"].font = Font(size=14, bold=True)
ws5["A2"] = "Registros reales en disco vs. los que la app devolvía por mes, y registros que se perdían por pertenecer a otro empleado."
ws5["A2"].font = Font(size=9, italic=True, color="FF666666")
ws5.append([])
ws5.append(["Mes", "Registros reales", "La app mostraba", "Ocultos (de otro empleado)", "Horas ocultas"])
hdr5 = ws5.max_row
for c in range(1, 6):
    cell = ws5.cell(row=hdr5, column=c)
    cell.fill = AZUL
    cell.font = BLANCO_BOLD
    cell.alignment = Alignment(horizontal="center", wrap_text=True)
for m in data["resumenGlobal"]:
    ws5.append([m["mes"], m["reales"], m["mostradosAntes"], m["ocultos"], m["horasOcultas"]])
    ws5.cell(row=ws5.max_row, column=5).number_format = "0.00"
ws5.append(["TOTAL", sum(m["reales"] for m in data["resumenGlobal"]),
            sum(m["mostradosAntes"] for m in data["resumenGlobal"]),
            sum(m["ocultos"] for m in data["resumenGlobal"]),
            sum(m["horasOcultas"] for m in data["resumenGlobal"])])
r = ws5.max_row
for c in range(1, 6):
    ws5.cell(row=r, column=c).fill = GRIS
    ws5.cell(row=r, column=c).font = BOLD
ws5.cell(row=r, column=5).number_format = "0.00"
for i, w in enumerate([16, 18, 16, 26, 14], start=1):
    ws5.column_dimensions[get_column_letter(i)].width = w

# ---- Hoja 6: empleados afectados en todo el año
ws6 = wb.create_sheet("Ocultos año 2026")
ws6["A1"] = "EMPLEADOS CUYOS REGISTROS NO MOSTRABA LA APP (MARZO - OCTUBRE 2026)"
ws6["A1"].font = Font(size=14, bold=True)
ws6["A2"] = ("Solo casos entre empleados distintos (el registro se perdía por pertenecer a otra persona). "
             "Los duplicados del mismo empleado y los nombres alternos de una misma persona se excluyeron.")
ws6["A2"].font = Font(size=9, italic=True, color="FF666666")
ws6.append([])
ws6.append(["Empleado", "Sucursal", "Registros ocultos", "Horas ocultas", "Meses afectados", "Registro que sí aparecía (de otro empleado)"])
hdr6 = ws6.max_row
for c in range(1, 7):
    cell = ws6.cell(row=hdr6, column=c)
    cell.fill = AZUL
    cell.font = BLANCO_BOLD
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

anual = {}
for mes, mdata in data["months"].items():
    for h in mdata["ocultos"]:
        e = anual.setdefault(nombre(h), {"suc": h.get("branch") or "—", "n": 0, "h": 0.0,
                                         "meses": set(), "abs": set()})
        e["n"] += 1
        e["h"] += float(h.get("totalHours") or 0)
        e["meses"].add(mes[:3])
        e["abs"].add(h.get("absorbidoPor") or "?")
for emp, e in sorted(anual.items(), key=lambda kv: -kv[1]["h"]):
    ws6.append([emp, e["suc"], e["n"], e["h"], ", ".join(sorted(e["meses"])), ", ".join(sorted(e["abs"]))])
    r = ws6.max_row
    ws6.cell(row=r, column=4).number_format = "0.00"
    ws6.cell(row=r, column=4).font = BOLD
ws6.append(["TOTAL", "", sum(e["n"] for e in anual.values()), sum(e["h"] for e in anual.values()), "", ""])
r = ws6.max_row
for c in range(1, 7):
    ws6.cell(row=r, column=c).fill = GRIS
    ws6.cell(row=r, column=c).font = BOLD
ws6.cell(row=r, column=4).number_format = "0.00"
for i, w in enumerate([32, 14, 16, 13, 22, 42], start=1):
    ws6.column_dimensions[get_column_letter(i)].width = w
ws6.freeze_panes = f"A{hdr6 + 1}"

wb.save(OUT_PATH)

print("OK ->", OUT_PATH)
print("Hojas:", wb.sheetnames)
print("Nómina: %d empleados con horas aprobadas | total a pagar $%.2f" % (len(nomina), sum(x[10] for x in nomina)))
print("Detalle: %d registros | ocultos por el bug: %d (%.2f h)" % (len(records), len(ocultos), sum(float(h.get("totalHours") or 0) for h in ocultos)))
